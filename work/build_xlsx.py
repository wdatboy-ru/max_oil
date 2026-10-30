# Сборка FINAL_parameters.xlsx из таблиц FINAL_REPORT.md + чек-лист + все находки Аргона
import re, csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from checklist import ITEMS, scores

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__))) + '/'
md = open(ROOT + 'FINAL_REPORT.md', encoding='utf-8').read().splitlines()
LINK = re.compile(r'\[([^\]]+)\]\((file:///[^)]+)\)')

SHEETS = {  # заголовок раздела → имя листа
    '1. Ответ': 'Ответ',
    '2. Чек-лист входных данных (по мануалам) и статус': 'Чек-лист (мануалы)',
    '3.1 Ekofisk (Danian)': 'Ekofisk',
    '3.2 Ommelanden / Tor (Maastrichtian–Campanian)': 'Ommelanden-Tor',
    '3.3 Pliocene (Oosterhout, газ)': 'Pliocene',
    '4.1 SCAL': 'SCAL',
    '4.2 Трещины и DFN': 'Трещины-DFN',
    '4.3 ГДИ и выводы': 'ГДИ',
    '4.4 Давления по годам': 'Давления по годам',
    '4.5 PVT': 'PVT',
    '4.6 Контакты': 'Контакты',
    '4.7 Газ плиоцена и запасы': 'Плиоцен и запасы',
    '4.8 Каротаж и методики': 'Каротаж',
    '4.9 Литология, фации, референсные значения': 'Литология',
    '5. Пробелы и альтернативы': 'Пробелы',
    '6. Сверка с твоим прежним файлом (PRJ)': 'Сверка с PRJ',
    '7. Ошибки отчётов Аргона, найденные при проверке': 'Ошибки Аргона',
    '8. Итоговые значения для модели': 'Итог для модели',
}

wb = Workbook(); wb.remove(wb.active)
head_font = Font(bold=True, color='FFFFFF'); head_fill = PatternFill('solid', fgColor='305496')
wrap = Alignment(wrap_text=True, vertical='top')


def plain(s):
    return LINK.sub(lambda m: m.group(1), s).replace('**', '').replace('`', '').strip()


def add_table(ws, rows, start):
    """rows: список списков ячеек markdown; добавляет столбец «Ссылка» с первой ссылкой строки."""
    r0 = start
    for i, cells in enumerate(rows):
        urls = [u for c in cells for _, u in LINK.findall(c)]
        vals = [plain(c) for c in cells] + (['Ссылка'] if i == 0 else [urls[0] if urls else ''])
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r0, column=j, value=v)
            cell.alignment = wrap
            if i == 0:
                cell.font = head_font; cell.fill = head_fill
            elif j == len(vals) and v:
                cell.hyperlink = v; cell.value = 'открыть'; cell.font = Font(color='0563C1', underline='single')
        r0 += 1
    return r0 + 1


section, tables, cur = None, {}, None
for line in md:
    h = re.match(r'^#{2,3} (.+)$', line)
    if h:
        section = h.group(1).strip(); cur = None; continue
    if line.startswith('|'):
        cells = [c for c in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r'\s*:?-{3,}:?\s*', c) for c in cells):
            continue
        if cur is None:
            cur = []; tables.setdefault(section, []).append(cur)
        cur.append([c.strip() for c in cells])
    else:
        cur = None

for sec, name in SHEETS.items():
    if sec not in tables:
        print('нет таблиц в разделе', sec); continue
    ws = wb.create_sheet(name[:31]); r = 1
    ws.cell(row=r, column=1, value=sec).font = Font(bold=True, size=12); r += 2
    for t in tables[sec]:
        r = add_table(ws, t, r)
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = [28, 55, 22, 40, 18, 40, 12, 12][col - 1]
    ws.freeze_panes = 'A4'

# чек-лист с баллами
ws = wb.create_sheet('Баллы полноты')
comp, rel, groups = scores()
num = sum(i[3] * i[2] * i[4] for i in ITEMS); den = sum(i[2] * i[4] for i in ITEMS)
ws.append(['Группа', 'Параметр', 'Полнота (0–1)', 'Надёжность (0–1)', 'Вес', 'Комментарий'])
for c in ws[1]:
    c.font = head_font; c.fill = head_fill
for it in ITEMS:
    ws.append(list(it))
ws.append([]); ws.append(['ИТОГО', 'полнота (взвеш.)', round(comp, 3)]); ws.append(['', 'надёжность имеющихся (взвеш.)', round(num / den, 3)])
for col, w in zip('ABCDEF', (16, 48, 14, 16, 6, 70)):
    ws.column_dimensions[col].width = w

# все находки Аргона
ws = wb.create_sheet('Находки Аргона (все)')
ws.append(['№ отчёта', 'Документ (relative)', 'Часть', 'Тема', 'Стр. оригинала', 'Скважины', 'НЕЧИТАЕМО', 'Текст находки', 'Ссылка'])
for c in ws[1]:
    c.font = head_font; c.fill = head_fill
rows = list(csv.DictReader(open(ROOT + 'work/all_findings.csv', encoding='utf-8-sig')))
for x in rows:
    txt = re.sub(r'\s+', ' ', x['text'])[:3000]
    ws.append([x['rid'], x['relative'], x['part'], x['topic_name'], x['orig_page'], x['wells'],
               'да' if x['unreadable'] == 'True' else '', txt, 'открыть'])
    c = ws.cell(row=ws.max_row, column=9); c.hyperlink = x['link']; c.font = Font(color='0563C1', underline='single')
for col, w in zip('ABCDEFGHI', (8, 50, 12, 18, 8, 16, 10, 120, 9)):
    ws.column_dimensions[col].width = w
ws.auto_filter.ref = f'A1:I{ws.max_row}'; ws.freeze_panes = 'A2'

wb.save(ROOT + 'FINAL_parameters.xlsx')
print('sheets:', wb.sheetnames, 'findings:', len(rows), f'полнота {comp:.2f}, надёжность {num/den:.2f}')
